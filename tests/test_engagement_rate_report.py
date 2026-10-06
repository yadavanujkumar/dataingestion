from __future__ import annotations

from pathlib import Path
import openpyxl
import pytest

from howlplatform.deliverables.contract import get_generator
from howlplatform.platform.canonical.schemas import PaidRowV2, OrganicRowV1
from howlplatform.platform.manifest.loader import load_manifest

REPO_ROOT = Path(__file__).parent.parent


def test_engagement_rate_report_ratio_of_sums(tmp_path):
    gen_cls = get_generator("engagement_rate_report")
    assert gen_cls is not None
    generator = gen_cls()

    manifest = load_manifest(REPO_ROOT / "manifests" / "ampere.json")

    # Day 1: 100,000 impressions, 15,000 engagements (15%)
    # Day 2: 10,000 impressions, 500 engagements (5%)
    # Ratio of sums = (15,000 + 500) / (100,000 + 10,000) = 15,500 / 110,000 = 14.09%
    # Erroneous average of percentages would be (15% + 5%) / 2 = 10.00%
    paid_rows = [
        PaidRowV2(
            brand_id="ampere",
            date="2026-08-01",
            channel="instagram",
            campaign="Ampere_IG_1",
            spend=5000.0,
            impressions=100000,
            clicks=2000,
            engagements=15000,
        ),
        PaidRowV2(
            brand_id="ampere",
            date="2026-08-02",
            channel="instagram",
            campaign="Ampere_IG_2",
            spend=500.0,
            impressions=10000,
            clicks=100,
            engagements=500,
        ),
    ]

    organic_rows = [
        OrganicRowV1(
            brand_id="ampere",
            date="2026-08-01",
            channel="instagram",
            post_id="p1",
            impressions=5000,
            engagements=250,
        )
    ]

    out_files = generator.build(
        brand=manifest,
        period="2026-08",
        data={"paid": paid_rows, "organic": organic_rows},
        out_dir=tmp_path,
    )

    assert len(out_files) == 1
    wb = openpyxl.load_workbook(out_files[0])
    assert "Engagement Rates" in wb.sheetnames
    assert "Methodology" in wb.sheetnames

    ws = wb["Engagement Rates"]
    # Check Instagram Paid row:
    # 110,000 impressions, 15,500 engagements -> 14.09%
    ig_row_values = [list(r) for r in ws.iter_rows(values_only=True) if r[0] == "Instagram"]
    assert len(ig_row_values) >= 1
    # Check that rate string is '14.09%' (ratio of sums, NOT 10.00%)
    assert ig_row_values[0][3] == "14.09%"
