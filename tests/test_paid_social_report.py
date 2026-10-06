from __future__ import annotations

from pathlib import Path
import openpyxl
import pytest

from howlplatform.deliverables.contract import get_generator
from howlplatform.platform.canonical.schemas import PaidRowV2
from howlplatform.platform.manifest.loader import load_manifest

REPO_ROOT = Path(__file__).parent.parent


def test_paid_social_report_generator_build(tmp_path):
    gen_cls = get_generator("paid_social_report")
    assert gen_cls is not None
    generator = gen_cls()

    manifest = load_manifest(REPO_ROOT / "manifests" / "ampere.json")
    rows = [
        PaidRowV2(
            brand_id="ampere",
            date="2026-08-01",
            channel="meta_ads",
            campaign="Ampere_Magnus_Awareness",
            spend=15000.0,
            impressions=200000,
            clicks=4000,
            engagements=30000,
            conversions=25.0,
        )
    ]

    out_files = generator.build(
        brand=manifest,
        period="2026-08",
        data={"paid": rows},
        out_dir=tmp_path,
    )

    assert len(out_files) == 1
    wb = openpyxl.load_workbook(out_files[0])
    assert "Summary" in wb.sheetnames
    assert "Notes" in wb.sheetnames
